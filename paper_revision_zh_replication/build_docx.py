"""Build the APA-formatted Word manuscript from the Markdown paper.

The Markdown is parsed by deterministic line scanning rather than by regular
expressions with open-ended lookaheads: an earlier version let one such pattern
swallow the whole document and duplicate it into the output.  This module
extracts the front matter, body, references, appendix, and float blocks as
disjoint slices, renders each with python-docx, and then imposes the submission
formatting: Times New Roman 12 pt, 1.5 line spacing, 1-inch margins, a page
number in the header, a double-spaced title page, APA heading levels, a
hanging-indent reference list, and tables and figures placed after the
reference list with each item starting its own page.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Final_Paper_EN_Chinese_Replication.md"
OUTPUT = ROOT / "Final_Paper_EN_Chinese_Replication.docx"

FONT = "Times New Roman"
SIZE = Pt(12)
TABLE_SIZE = Pt(9)
LINE_SPACING = 1.5

# APA 7 table geometry, shared by Table 1 (9 columns) and Table A1 (7 columns).
WIDTHS_9 = [Inches(1.15), Inches(0.42), Inches(0.45), Inches(0.45), Inches(0.42),
            Inches(0.62), Inches(0.62), Inches(0.85), Inches(0.62)]
WIDTHS_7 = [Inches(1.45), Inches(0.55), Inches(0.50), Inches(0.80), Inches(1.30),
            Inches(1.30), Inches(0.85)]
FIGURE_WIDTH = Inches(6.0)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


# --------------------------------------------------------------------------- #
# Markdown source
# --------------------------------------------------------------------------- #

def parse_blocks(lines: list[str], *, body_out: list[str]) -> tuple[list[list[str]], list[str]]:
    """Split markdown lines into float blocks and everything else.

    A float begins at a column-0 "**Table n**" or "**Figure n**" label and ends
    at the first blank line that is not followed by a table row.  The blank-line
    rule matters: a float cannot run to the next label, because that would
    swallow whole Results subsections, and it must include a markdown table plus
    its note, which are separated from the label by blank lines.
    """
    floats: list[list[str]] = []
    body: list[str] = []
    block: list[str] | None = None
    for i, line in enumerate(lines):
        if re.match(r"^\*\*?(?:Table|Figure) ", line):
            if block:
                floats.append(block)
            block = [line]
            continue
        if block is not None:
            if line.strip():
                block.append(line)
            else:
                # Keep the block open while a markdown table or its note is
                # still ahead.  Table rows can be several blank lines below the
                # label because the title sits between them.
                ahead = lines[i + 1:]
                still_float = False
                for candidate in ahead:
                    stripped = candidate.lstrip()
                    if stripped.startswith("|"):
                        still_float = True
                        break
                    if candidate.startswith("#") or re.match(r"^\*\*?(?:Table|Figure) ", candidate):
                        break
                if still_float:
                    block.append(line)
                else:
                    floats.append(block)
                    block = None
            continue
        body.append(line)
    if block:
        floats.append(block)
    body_out.extend(body)
    return floats, body


class Paper:
    """Disjoint slices of the source manuscript."""

    def __init__(self, md: str) -> None:
        lines = md.split("\n")

        def index_of(prefix: str, start: int = 0) -> int:
            for i in range(start, len(lines)):
                if lines[i].startswith(prefix):
                    return i
            raise ValueError(f"section not found: {prefix!r}")

        self.title = lines[0][2:].strip()
        # Front matter order is title, project line, then the byline block
        # (author, affiliation, mentor).  Locate them by content rather than by
        # fixed offsets so that adding a line cannot silently empty the byline.
        self.project = next(
            (ln.strip().strip("*") for ln in lines[1:12] if ln.startswith("*Project")), ""
        )
        author_at = next(
            (i for i in range(1, 12) if lines[i].strip() and not lines[i].startswith(("*", "#"))), 1
        )
        byline_end = author_at
        while byline_end < len(lines) and lines[byline_end].strip():
            byline_end += 1
        self.byline = [ln.strip() for ln in lines[author_at:byline_end] if ln.strip()]

        abs_at = index_of("## Abstract")
        kw_at = index_of("*Keywords:")
        self.abstract = " ".join(ln.strip() for ln in lines[abs_at + 1:kw_at] if ln.strip())
        self.keywords = lines[kw_at].split(":", 1)[1].strip()

        body_at = index_of("## Related Work")
        refs_at = index_of("## References")
        app_at = index_of("## Appendix A")

        body_lines = lines[body_at:refs_at]
        self.references = [ln.strip() for ln in lines[refs_at + 1:app_at] if ln.strip()]

        # Body and floats, split by the shared parser.
        self.body = []
        self.floats, _ = parse_blocks(body_lines, body_out=self.body)

        # The appendix carries one float (its full-cohort table) plus prose.
        self.appendix_body = []
        self.appendix_floats, _ = parse_blocks(lines[app_at:], body_out=self.appendix_body)


# --------------------------------------------------------------------------- #
# Rendering primitives
# --------------------------------------------------------------------------- #

def new_document() -> docx.Document:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)

    section = document.sections[0]
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(1))
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)
    run.font.name = FONT
    run.font.size = SIZE
    return document


def configure(paragraph, *, spacing: float = LINE_SPACING, indent: float = 0.0,
              hanging: bool = False, break_before: bool = False,
              align=WD_ALIGN_PARAGRAPH.LEFT, space_before: float = 0.0,
              space_after: float = 0.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = spacing
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.widow_control = True
    fmt.page_break_before = break_before
    if hanging:
        fmt.left_indent = Inches(0.5)
        fmt.first_line_indent = Inches(-0.5)
    else:
        fmt.left_indent = Inches(0)
        fmt.first_line_indent = Inches(indent)
    paragraph.alignment = align


def add_runs(paragraph, text: str, *, size=SIZE, bold: bool = False,
             italic: bool = False) -> None:
    """Render inline **bold**, *italic*, and `code` spans."""
    for token in INLINE.split(text):
        if not token:
            continue
        is_bold, is_italic, content = bold, italic, token
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            is_bold, content = True, token[2:-2]
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            is_italic, content = True, token[1:-1]
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            content = token[1:-1]
        run = paragraph.add_run(content)
        run.font.name = FONT
        run.font.size = size
        run.font.bold = is_bold
        run.font.italic = is_italic
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_markdown_table(document, rows: list[list[str]], widths) -> None:
    if len(widths) != len(rows[0]):
        raise ValueError(
            f"table geometry mismatch: {len(rows[0])} columns but {len(widths)} widths"
        )
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        for existing in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(existing)
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(w.inches for w in widths) * 1440)))
    tbl_pr.append(tbl_w)
    for existing in tbl.findall(qn("w:tblGrid")):
        tbl.remove(existing)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(width.inches * 1440)))
        grid.append(column)
    tbl_pr.addnext(grid)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0"); element.set(qn("w:color"), "000000")
        borders.append(element)
    for edge in ("left", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "none"); element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0"); element.set(qn("w:color"), "auto")
        borders.append(element)
    tbl_pr.append(borders)

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true")
    tr_pr.append(header)

    for r, row in enumerate(rows):
        for c, (cell, text) in enumerate(zip(table.rows[r].cells, row)):
            cell.width = widths[c]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(widths[c].inches * 1440)))
            tc_pr.append(tc_w)
            paragraph = cell.paragraphs[0]
            configure(paragraph, spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT,
                      space_before=2, space_after=2)
            add_runs(paragraph, text, size=TABLE_SIZE, bold=(r == 0))
            if r == 0:
                tc_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:space"), "0"); bottom.set(qn("w:color"), "000000")
                tc_borders.append(bottom)
                tc_pr.append(tc_borders)


# --------------------------------------------------------------------------- #
# Block rendering
# --------------------------------------------------------------------------- #

def render_body(document, lines: list[str], *, first_break: bool = False) -> None:
    pending_break = first_break
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("#### "):
            paragraph = document.add_paragraph()
            configure(paragraph, break_before=pending_break, space_before=12)
            add_runs(paragraph, line[5:], bold=True, italic=True)
        elif line.startswith("### "):
            paragraph = document.add_paragraph()
            configure(paragraph, break_before=pending_break, space_before=12)
            add_runs(paragraph, line[4:], bold=True)
        elif line.startswith("## "):
            paragraph = document.add_paragraph()
            configure(paragraph, break_before=pending_break,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
            add_runs(paragraph, line[3:], bold=True)
        elif line.startswith("!["):
            match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            if match:
                path = (ROOT / match.group(1)).resolve()
                paragraph = document.add_paragraph()
                configure(paragraph, break_before=pending_break,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
                paragraph.add_run().add_picture(str(path), width=FIGURE_WIDTH)
        else:
            is_note = line.startswith("*Note.*")
            paragraph = document.add_paragraph()
            configure(paragraph, break_before=pending_break,
                      indent=0.0 if is_note else 0.5)
            add_runs(paragraph, line)
        pending_break = False


def render_float(document, block: list[str]) -> None:
    """A Table/Figure group: label, title, optional markdown table, image, note."""
    first = True
    i = 0
    while i < len(block):
        line = block[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(block) and block[i].lstrip().startswith("|"):
                cells = [c.strip() for c in block[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                widths = WIDTHS_9 if len(rows[0]) == 9 else WIDTHS_7
                add_markdown_table(document, rows, widths)
            continue
        if line.startswith("!["):
            match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            if match:
                path = (ROOT / match.group(1)).resolve()
                paragraph = document.add_paragraph()
                configure(paragraph, break_before=first,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
                paragraph.add_run().add_picture(str(path), width=FIGURE_WIDTH)
                first = False
            i += 1
            continue
        paragraph = document.add_paragraph()
        label = re.match(r"^\*\*?(Table|Figure) (A?\d+)\*\*?$", line)
        if label:
            configure(paragraph, break_before=first, space_after=6)
            add_runs(paragraph, f"{label.group(1)} {label.group(2)}", bold=True)
        else:
            is_note = line.startswith("*Note.*")
            italic_title = (line.startswith("*") and line.endswith("*")
                            and not is_note)
            configure(paragraph, break_before=first, indent=0.0)
            add_runs(paragraph, line.strip("*") if italic_title else line,
                     italic=italic_title or is_note)
        first = False
        i += 1


# --------------------------------------------------------------------------- #
# Assembly
# --------------------------------------------------------------------------- #

def build() -> Path:
    paper = Paper(SOURCE.read_text(encoding="utf-8"))
    document = new_document()

    # Title page (double-spaced per APA 7).
    paragraph = document.add_paragraph()
    configure(paragraph, spacing=2.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(paragraph, paper.title, bold=True)
    for entry in paper.byline:
        paragraph = document.add_paragraph()
        configure(paragraph, spacing=2.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(paragraph, entry)
    paragraph = document.add_paragraph()
    configure(paragraph, spacing=2.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(paragraph, paper.project, italic=True)

    # Abstract page.
    paragraph = document.add_paragraph()
    configure(paragraph, break_before=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(paragraph, "Abstract", bold=True)
    paragraph = document.add_paragraph()
    configure(paragraph, indent=0.5)
    add_runs(paragraph, paper.abstract)
    paragraph = document.add_paragraph()
    configure(paragraph, indent=0.5)
    add_runs(paragraph, f"Keywords: {paper.keywords}", italic=True)

    # Body, references, appendix, then the table and figure blocks.
    render_body(document, paper.body, first_break=True)

    paragraph = document.add_paragraph()
    configure(paragraph, break_before=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(paragraph, "References", bold=True)
    for entry in paper.references:
        paragraph = document.add_paragraph()
        configure(paragraph, hanging=True)
        add_runs(paragraph, entry)

    render_body(document, paper.appendix_body, first_break=True)
    for block in paper.appendix_floats:
        render_float(document, block)

    for block in paper.floats:
        render_float(document, block)

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
    print(f"size: {path.stat().st_size:,} bytes")
    sys.exit(0)
